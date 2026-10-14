import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

F="Arial"; NAVY="1C3359"; ALT="F5F8FC"
H=Font(name=F,bold=True,color="FFFFFF"); HF=PatternFill("solid",fgColor=NAVY)
B=Font(name=F); BB=Font(name=F,bold=True); BLUE=Font(name=F,color="0000FF"); TITLE=Font(name=F,bold=True,size=14,color=NAVY)
GREY=Font(name=F,color="5A6472",italic=True); YEL=PatternFill("solid",fgColor="FFFF00")
TT='"TT$"#,##0.00;("TT$"#,##0.00);"-"'; PCT='0.0%'; DATE='dd mmm yyyy'; MON='mmm yyyy'
N=150

wb=Workbook()
# ---------- Settings ----------
s=wb.active; s.title="Settings"
s["A1"]="Settings"; s["A1"].font=TITLE
s["A2"]="Blue text = numbers you can change. Yellow = fill in. Black = formulas, don't type over them."; s["A2"].font=GREY
s["A4"]="How each payment is split"; s["A4"].font=BB
rows=[("Tax & NIS savings",0.20),("Business costs",0.15),("Emergency fund",0.15)]
for i,(k,v) in enumerate(rows,start=5):
    s[f"A{i}"]=k; s[f"B{i}"]=v; s[f"B{i}"].font=BLUE; s[f"B{i}"].number_format=PCT
s["A8"]="Pay yourself"; s["B8"]="=1-SUM(B5:B7)"; s["B8"].number_format=PCT
s["A9"]="Total (must be 100%)"; s["B9"]="=SUM(B5:B8)"; s["B9"].number_format=PCT; s["A9"].font=BB
s["A11"]="Goals"; s["A11"].font=BB
s["A12"]="Emergency fund target"; s["B12"]=None; s["B12"].fill=YEL
s["C12"]="About 3 months of your bills"; s["C12"].font=GREY
s["A13"]="PC build budget"; s["B13"].fill=YEL; s["C13"]="Your total quote for the PC"; s["C13"].font=GREY
s["A14"]="Claude Max cost per month"; s["B14"]=680; s["B14"].font=BLUE
s["C14"]="About US$100 a month; check the current price and exchange rate"; s["C14"].font=GREY
s["B14"].comment=Comment("Source: approximate, Claude Max 5x at about US$100/month converted at about TT$6.8 per US$. Check current price.","Portside Digital")
for r in (12,13,14): s[f"B{r}"].number_format=TT
s["A16"]="Saved toward PC so far"; s["B16"]=0; s["B16"].font=BLUE; s["B16"].number_format=TT
s["C16"]="Update this when you move money into PC savings"; s["C16"].font=GREY
s.column_dimensions["A"].width=30; s.column_dimensions["B"].width=16; s.column_dimensions["C"].width=58

# ---------- Income ----------
inc=wb.create_sheet("Income")
hdr=["Date","Client","What for","Amount (TT$)","Status","Payment method","Notes",
     "Tax & NIS","Business costs","Emergency fund","Pay yourself","Month"]
inc.append(hdr)
inc.append([dt.date(2026,10,13),"Example Boutique","Website deposit (50%)",8000,"Paid","Bank transfer","Example row: replace or delete"])
for r in range(2,N+2):
    inc[f"H{r}"]=f'=IF($E{r}="Paid",$D{r}*Settings!$B$5,"")'
    inc[f"I{r}"]=f'=IF($E{r}="Paid",$D{r}*Settings!$B$6,"")'
    inc[f"J{r}"]=f'=IF($E{r}="Paid",$D{r}*Settings!$B$7,"")'
    inc[f"K{r}"]=f'=IF($E{r}="Paid",$D{r}*Settings!$B$8,"")'
    inc[f"L{r}"]=f'=IF($A{r}="","",DATE(YEAR($A{r}),MONTH($A{r}),1))'
    inc[f"A{r}"].number_format=DATE; inc[f"L{r}"].number_format=MON
    for c in "DHIJK": inc[f"{c}{r}"].number_format=TT
    for c in "ABCDEFG": inc[f"{c}{r}"].font=BLUE if c=="D" else B
dv=DataValidation(type="list",formula1='"Paid,Owed,Refunded"',allow_blank=True); inc.add_data_validation(dv); dv.add(f"E2:E{N+1}")
dv2=DataValidation(type="list",formula1='"Bank transfer,WiPay,Fygaro,Cash,Other"',allow_blank=True); inc.add_data_validation(dv2); dv2.add(f"F2:F{N+1}")
widths=[13,22,30,15,11,16,30,14,15,15,14,11]
for i,w in enumerate(widths): inc.column_dimensions[chr(65+i)].width=w

# ---------- Expenses ----------
ex=wb.create_sheet("Expenses")
ex.append(["Date","What","Category","Amount (TT$)","Receipt saved?","Notes","Month"])
ex.append([dt.date(2026,10,13),"Claude Pro plan","Claude",140,"Yes","Example row: replace or delete"])
for r in range(2,N+2):
    ex[f"G{r}"]=f'=IF($A{r}="","",DATE(YEAR($A{r}),MONTH($A{r}),1))'
    ex[f"A{r}"].number_format=DATE; ex[f"G{r}"].number_format=MON; ex[f"D{r}"].number_format=TT; ex[f"D{r}"].font=BLUE
dv3=DataValidation(type="list",formula1='"Ads,Claude,Domains & hosting,Equipment,Software,Other"',allow_blank=True); ex.add_data_validation(dv3); dv3.add(f"C2:C{N+1}")
dv4=DataValidation(type="list",formula1='"Yes,No"',allow_blank=True); ex.add_data_validation(dv4); dv4.add(f"E2:E{N+1}")
for i,w in enumerate([13,30,18,15,15,34,11]): ex.column_dimensions[chr(65+i)].width=w

# ---------- Summary ----------
sm=wb.create_sheet("Summary",0)
sm["A1"]="Portside Digital: Money Summary"; sm["A1"].font=TITLE
sm["A2"]="Everything here is calculated from the Income, Expenses and Settings tabs."; sm["A2"].font=GREY
items=[
 ("Totals",None),
 ("Total earned (paid)",'=SUMIFS(Income!D:D,Income!E:E,"Paid")'),
 ("Still owed by clients",'=SUMIFS(Income!D:D,Income!E:E,"Owed")'),
 ("Business spending",'=SUM(Expenses!D:D)'),
 ("Where the money goes",None),
 ("Tax & NIS set aside",'=SUM(Income!H:H)'),
 ("Business costs money",'=SUM(Income!I:I)'),
 ("Business money left after spending",'=B10-B7'),
 ("Emergency fund",'=SUM(Income!J:J)'),
 ("Pay yourself",'=SUM(Income!K:K)'),
 ("Goals",None),
 ("Months of Claude Max covered",'=IF(Settings!B14=0,0,B11/Settings!B14)'),
 ("Emergency fund progress",'=IF(N(Settings!B12)=0,0,B12/Settings!B12)'),
 ("Left to save for the PC",'=MAX(0,N(Settings!B13)-Settings!B16)'),
 ("PC savings progress",'=IF(N(Settings!B13)=0,0,Settings!B16/Settings!B13)'),
]
r=4
for k,f in items:
    sm[f"A{r}"]=k
    if f is None: sm[f"A{r}"].font=BB
    else:
        sm[f"B{r}"]=f
        sm[f"B{r}"].number_format = '0.0' if "Months" in k else (PCT if "progress" in k else TT)
    r+=1
# r now 19
sm["A21"]="By month"; sm["A21"].font=BB
mh=["Month","Earned (paid)","Still owed","Business spending","Tax & NIS set aside","Pay yourself"]
for i,h in enumerate(mh): sm.cell(row=22,column=i+1,value=h)
for i in range(12):
    rr=23+i
    sm[f"A{rr}"]= dt.date(2026,10,1) if i==0 else f"=EDATE(A{rr-1},1)"
    if i==0:
        sm[f"A{rr}"].font=BLUE; sm[f"A{rr}"].comment=Comment("First month of the table. Change it and the rest follow.","Portside Digital")
    sm[f"A{rr}"].number_format=MON
    sm[f"B{rr}"]=f'=SUMIFS(Income!$D:$D,Income!$L:$L,$A{rr},Income!$E:$E,"Paid")'
    sm[f"C{rr}"]=f'=SUMIFS(Income!$D:$D,Income!$L:$L,$A{rr},Income!$E:$E,"Owed")'
    sm[f"D{rr}"]=f'=SUMIFS(Expenses!$D:$D,Expenses!$G:$G,$A{rr})'
    sm[f"E{rr}"]=f'=SUMIFS(Income!$H:$H,Income!$L:$L,$A{rr})'
    sm[f"F{rr}"]=f'=SUMIFS(Income!$K:$K,Income!$L:$L,$A{rr})'
    for c in "BCDEF": sm[f"{c}{rr}"].number_format=TT
sm["A35"]="Total"; sm["A35"].font=BB
for c in "BCDEF":
    sm[f"{c}35"]=f"=SUM({c}23:{c}34)"; sm[f"{c}35"].number_format=TT; sm[f"{c}35"].font=BB
sm.column_dimensions["A"].width=34
for c in "BCDEF": sm.column_dimensions[c].width=19

# header styling
for ws,row in ((inc,1),(ex,1),(sm,22)):
    for cell in ws[row]:
        if cell.value is not None:
            cell.font=H; cell.fill=HF; cell.alignment=Alignment(vertical="center",wrap_text=True)
inc.freeze_panes="A2"; ex.freeze_panes="A2"
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.font and c.font.name!=F: c.font=Font(name=F,bold=c.font.bold,color=c.font.color,italic=c.font.italic,size=c.font.sz)
wb.save("money_tracker.xlsx")
